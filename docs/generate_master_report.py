import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def convert_report():
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    filepath = 'SYSTEM_COMPLIANCE_REPORT.md'
    if not os.path.exists(filepath):
        print("Error: SYSTEM_COMPLIANCE_REPORT.md not found.")
        return
            
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue

        # Headings
        if line.startswith('# '):
            h = doc.add_heading(line[2:], level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        
        # Lists
        elif line.startswith('* ') or line.startswith('- '):
            if '[x]' in line:
                p = doc.add_paragraph('☑ ' + line.replace('- [x]', '').replace('* [x]', '').strip())
            elif '[ ]' in line:
                p = doc.add_paragraph('☐ ' + line.replace('- [ ]', '').replace('* [ ]', '').strip())
            else:
                p = doc.add_paragraph(line[2:], style='List Bullet')
        elif re.match(r'^\d+\.', line):
            p = doc.add_paragraph(re.sub(r'^\d+\.\s*', '', line), style='List Number')
        
        # Horizontal Rules
        elif line == '---':
            doc.add_paragraph('________________________________________________________________________________', style='Normal')
        
        # Normal text with bold handling
        else:
            p = doc.add_paragraph()
            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

    output_path = 'System_Compliance_Report.docx'
    doc.save(output_path)
    print(f"Report saved to {output_path}")

if __name__ == "__main__":
    convert_report()
